from pathlib import Path
from io import BytesIO

from openpyxl import load_workbook
from docx import Document
from pypdf import PdfReader


# =========================================================
# TXT
# =========================================================

def read_txt(inbox, attachment):
    return inbox.read_text(attachment)


# =========================================================
# XLSX
# =========================================================

def read_xlsx(inbox, attachment):

    data = inbox.read_bytes(attachment)

    workbook = load_workbook(
        filename=BytesIO(data),
        data_only=True
    )

    lines = []

    for sheet in workbook.worksheets:

        lines.append(f"Sheet: {sheet.title}")

        for row in sheet.iter_rows(values_only=True):

            values = []

            for value in row:

                if value is not None:

                    value = str(value).strip()

                    if value:
                        values.append(value)

            if values:
                lines.append(" | ".join(values))

    return "\n".join(lines)


# =========================================================
# DOCX
# =========================================================

def read_docx(inbox, attachment):

    data = inbox.read_bytes(attachment)

    document = Document(BytesIO(data))

    lines = []

    # Paragraphs
    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            lines.append(text)

    # Tables
    for table in document.tables:

        for row in table.rows:

            values = []

            for cell in row.cells:

                text = cell.text.strip()

                if text:
                    values.append(text)

            if values:
                lines.append(" | ".join(values))

    return "\n".join(lines)


# =========================================================
# PDF - PyMuPDF TEXT EXTRACTION
# =========================================================

def read_pdf_text_pymupdf(data):

    import pymupdf

    document = pymupdf.open(
        stream=data,
        filetype="pdf"
    )

    pages = []

    for page in document:

        text = page.get_text("text")

        if text:
            pages.append(text)

    document.close()

    result = "\n".join(pages).strip()

    return result


# =========================================================
# PDF - OCR
# =========================================================

def read_pdf_ocr(data):

    import pymupdf
    import pytesseract

    from PIL import Image

    # Tesseract executable
    pytesseract.pytesseract.tesseract_cmd = (
        r"C:\Program Files\Tesseract-OCR\tesseract.exe"
    )

    document = pymupdf.open(
        stream=data,
        filetype="pdf"
    )

    pages = []

    for page_number, page in enumerate(document):

        print(
            f"OCR processing page "
            f"{page_number + 1}..."
        )

        # Render PDF page as high-resolution image
        matrix = pymupdf.Matrix(
            2.5,
            2.5
        )

        pixmap = page.get_pixmap(
            matrix=matrix,
            alpha=False
        )

        image_bytes = pixmap.tobytes(
            "png"
        )

        image = Image.open(
            BytesIO(image_bytes)
        )

        # OCR
        text = pytesseract.image_to_string(
            image,
            config="--psm 6"
        )

        if text:

            pages.append(
                text
            )

    document.close()

    result = "\n".join(
        pages
    ).strip()

    if not result:

        raise ValueError(
            "OCR produced no text"
        )

    return result


# =========================================================
# PDF - PYPDF FALLBACK
# =========================================================

def read_pdf_pypdf(data):

    reader = PdfReader(
        BytesIO(data)
    )

    normal_pages = []
    layout_pages = []

    for page in reader.pages:

        # Normal extraction
        try:

            text = page.extract_text()

        except Exception:

            text = None

        if text:
            normal_pages.append(text)

        # Layout extraction
        try:

            layout_text = page.extract_text(
                extraction_mode="layout"
            )

        except Exception:

            layout_text = None

        if layout_text:
            layout_pages.append(
                layout_text
            )

    normal_text = "\n".join(
        normal_pages
    ).strip()

    layout_text = "\n".join(
        layout_pages
    ).strip()

    if layout_text:
        return layout_text

    if normal_text:
        return normal_text

    raise ValueError(
        "PDF contains no extractable text"
    )


# =========================================================
# PDF - MAIN READER
# =========================================================

def read_pdf(inbox, attachment):

    data = inbox.read_bytes(
        attachment
    )

    # =====================================================
    # METHOD 1: PyMuPDF TEXT
    # =====================================================

    try:

        text = read_pdf_text_pymupdf(
            data
        )

        if text:

            print(
                f"PDF text extraction: "
                f"{attachment}"
            )

            return text

    except Exception as error:

        print(
            f"PyMuPDF text failed: "
            f"{attachment} -> {error}"
        )


    # =====================================================
    # METHOD 2: OCR
    # =====================================================

    try:

        print(
            f"Image-only PDF detected. "
            f"Starting OCR: {attachment}"
        )

        text = read_pdf_ocr(
            data
        )

        if text:

            print(
                f"OCR successful: "
                f"{attachment}"
            )

            return text

    except Exception as error:

        print(
            f"OCR failed: "
            f"{attachment} -> {error}"
        )


    # =====================================================
    # METHOD 3: pypdf
    # =====================================================

    try:

        text = read_pdf_pypdf(
            data
        )

        if text:

            print(
                f"pypdf extraction: "
                f"{attachment}"
            )

            return text

    except Exception as error:

        print(
            f"pypdf failed: "
            f"{attachment} -> {error}"
        )


    # =====================================================
    # ALL METHODS FAILED
    # =====================================================

    raise ValueError(
        "PDF could not be read by "
        "PyMuPDF, OCR, or pypdf"
    )


# =========================================================
# MAIN DOCUMENT READER
# =========================================================

def read_document(inbox, attachment):

    extension = Path(
        attachment
    ).suffix.lower()

    try:

        # TXT
        if extension == ".txt":

            return read_txt(
                inbox,
                attachment
            )

        # XLSX
        if extension == ".xlsx":

            return read_xlsx(
                inbox,
                attachment
            )

        # DOCX
        if extension == ".docx":

            return read_docx(
                inbox,
                attachment
            )

        # PDF
        if extension == ".pdf":

            return read_pdf(
                inbox,
                attachment
            )

        # Unsupported
        raise ValueError(
            f"Unsupported document format: "
            f"{extension}"
        )

    except Exception as error:

        raise ValueError(
            f"Document reading failed: "
            f"{attachment} -> {error}"
        )