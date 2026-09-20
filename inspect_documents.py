import sys
from pathlib import Path
from io import BytesIO

from openpyxl import load_workbook
from docx import Document

BUNDLE_PATH = Path(r"C:\for\sdoc-hackathon-bundle")
sys.path.insert(0, str(BUNDLE_PATH))

from loader import Inbox


TARGET_EMAILS = [
    "email_097",
    "email_107",
    "email_243",
    "email_291",
    "email_300",
    "email_302",
    "email_435",
    "email_496",
]


def inspect_xlsx(inbox, attachment):
    print("\n--- EXCEL ---")

    data = inbox.read_bytes(attachment)
    workbook = load_workbook(
        filename=BytesIO(data),
        data_only=True
    )

    for sheet in workbook.worksheets:
        print(f"\n[SHEET: {sheet.title}]")

        for row in sheet.iter_rows():
            values = []

            for cell in row:
                if cell.value is not None:
                    values.append(
                        f"{cell.coordinate}={cell.value}"
                    )

            if values:
                print(" | ".join(values))


def inspect_docx(inbox, attachment):
    print("\n--- WORD ---")

    data = inbox.read_bytes(attachment)
    document = Document(BytesIO(data))

    print("\n[PARAGRAPHS]")

    for i, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()

        if text:
            print(f"P{i}: {text}")

    print("\n[TABLES]")

    for table_no, table in enumerate(document.tables):
        print(f"\nTABLE {table_no}")

        for row_no, row in enumerate(table.rows):
            values = []

            for cell_no, cell in enumerate(row.cells):
                text = cell.text.strip()

                values.append(
                    f"C{cell_no}={text}"
                )

            print(
                f"ROW {row_no}: "
                + " | ".join(values)
            )


def inspect_file(inbox, attachment):
    extension = Path(attachment).suffix.lower()

    print("\n" + "=" * 100)
    print(f"FILE: {attachment}")
    print("=" * 100)

    if extension == ".xlsx":
        inspect_xlsx(inbox, attachment)

    elif extension == ".docx":
        inspect_docx(inbox, attachment)

    else:
        print(f"Unsupported for inspection: {extension}")


def main():
    inbox = Inbox(str(BUNDLE_PATH))

    for email in inbox.emails():

        email_id = email["email_id"]

        if email_id not in TARGET_EMAILS:
            continue

        print("\n\n")
        print("#" * 100)
        print(f"# {email_id}")
        print("#" * 100)

        for attachment in email.get("attachments", []):
            extension = Path(attachment).suffix.lower()

            if extension in [".xlsx", ".docx"]:
                inspect_file(inbox, attachment)


if __name__ == "__main__":
    main()